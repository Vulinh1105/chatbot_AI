"""
parsing.py
----------
T14 - Doc parsing MVP.

Đọc toàn bộ file .doc / .docx trong thư mục backend/data/documents/,
trích xuất text thô, trả về dict {filename: text}.

- .docx: đọc trực tiếp bằng python-docx (không cần cài thêm phần mềm ngoài).
- .doc  (định dạng cũ): python-docx KHÔNG đọc được. Cách xử lý MVP:
    1) Nếu máy có LibreOffice (soffice), tự động convert .doc -> .docx rồi đọc.
    2) Nếu không có, ghi log cảnh báo và bỏ qua file đó (thay vì crash cả pipeline).

Cách chạy thử độc lập:
    python parsing.py
"""

from __future__ import annotations

import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document

logging.basicConfig(level=logging.INFO, format="[%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

# Mặc định trỏ tới backend/data/documents/ (tính từ vị trí file này:
# backend/app/doc_processing/parsing.py -> backend/data/documents/)
DEFAULT_DOCS_DIR = Path(__file__).resolve().parents[2] / "data" / "documents"

SUPPORTED_EXTENSIONS = {".docx", ".doc"}


def _read_docx(file_path: Path) -> str:
    """Đọc text từ file .docx (bao gồm cả text trong bảng)."""
    doc = Document(file_path)

    parts: list[str] = []

    # Đoạn văn bản thường
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())

    # Text trong bảng (nếu có)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    return "\n".join(parts)


def _convert_doc_to_docx(file_path: Path, tmp_dir: Path) -> Path | None:
    """
    Convert file .doc (định dạng cũ) sang .docx bằng LibreOffice (soffice),
    nếu LibreOffice có sẵn trên máy. Trả về None nếu không convert được.
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None

    try:
        subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                str(tmp_dir),
                str(file_path),
            ],
            check=True,
            capture_output=True,
            timeout=60,
        )
        converted = tmp_dir / (file_path.stem + ".docx")
        return converted if converted.exists() else None
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        logger.warning("Không convert được %s sang .docx: %s", file_path.name, e)
        return None


def parse_document(file_path: Path) -> str | None:
    """Đọc 1 file .doc/.docx, trả về text hoặc None nếu thất bại."""
    suffix = file_path.suffix.lower()

    if suffix == ".docx":
        try:
            return _read_docx(file_path)
        except Exception as e:
            logger.error("Lỗi khi đọc %s: %s", file_path.name, e)
            return None

    if suffix == ".doc":
        with tempfile.TemporaryDirectory() as tmp:
            converted = _convert_doc_to_docx(file_path, Path(tmp))
            if converted is None:
                logger.warning(
                    "Bỏ qua %s: đây là định dạng .doc cũ và không tìm thấy "
                    "LibreOffice (soffice) trên máy để convert. "
                    "Cài LibreOffice hoặc tự convert file này sang .docx.",
                    file_path.name,
                )
                return None
            try:
                return _read_docx(converted)
            except Exception as e:
                logger.error("Lỗi khi đọc %s (sau khi convert): %s", file_path.name, e)
                return None

    logger.warning("Bỏ qua %s: định dạng không được hỗ trợ.", file_path.name)
    return None


def parse_documents_in_dir(docs_dir: Path | str = DEFAULT_DOCS_DIR) -> dict[str, str]:
    """
    Duyệt toàn bộ file .doc/.docx trong docs_dir, trả về:
        {filename: raw_text}
    Các file lỗi hoặc không đọc được sẽ bị bỏ qua (đã log cảnh báo).
    """
    docs_dir = Path(docs_dir)
    if not docs_dir.exists():
        logger.error("Thư mục không tồn tại: %s", docs_dir)
        return {}

    results: dict[str, str] = {}

    files = sorted(
        p for p in docs_dir.iterdir()
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    )

    if not files:
        logger.warning("Không tìm thấy file .doc/.docx nào trong %s", docs_dir)
        return {}

    for file_path in files:
        logger.info("Đang xử lý: %s", file_path.name)
        text = parse_document(file_path)
        if text:
            results[file_path.name] = text
        else:
            logger.warning("Không lấy được text từ %s", file_path.name)

    logger.info("Hoàn tất: đọc thành công %d/%d file.", len(results), len(files))
    return results


if __name__ == "__main__":
    docs = parse_documents_in_dir()
    for name, text in docs.items():
        preview = text[:200].replace("\n", " ")
        print(f"\n--- {name} ({len(text)} ký tự) ---")
        print(preview + ("..." if len(text) > 200 else ""))